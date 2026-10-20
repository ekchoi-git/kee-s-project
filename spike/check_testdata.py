"""ChatGPT가 만든 테스트 파일의 '정답' 시트가 실제 데이터와 맞는지 직접 다시 계산해 확인."""
import sys, collections, datetime as dt
import openpyxl

def load(f):
    wb = openpyxl.load_workbook(f, read_only=True)
    data = list(wb["더미 데이터"].iter_rows(min_row=2, values_only=True))
    key = {r[0]: r[1] for r in wb["검증 정답"].iter_rows(values_only=True) if r[0] and isinstance(r[1], int)}
    agg = {}
    for r in wb["검증 정답"].iter_rows(values_only=True):
        if isinstance(r[1], int) and 2000 < r[1] < 2100 and isinstance(r[2], int):
            agg[(r[0], r[1])] = (r[2], r[3], r[6])
    return data, key, agg

def check(f):
    data, key, agg = load(f)
    num = lambda v: isinstance(v, (int, float))
    to_num = lambda v: v if num(v) else (float(v.replace(",", "")) if isinstance(v, str) and v.replace(",", "").lstrip("-").replace(".", "").isdigit() else None)
    got = {
        "데이터 행 수": len(data),
        "숫자 금액 합계": sum(r[3] for r in data if num(r[3])),
        "문자 변환 후 금액 합계": sum(to_num(r[3]) or 0 for r in data),
        "회사명 누락": sum(r[5] in (None, "") for r in data),
        "금액 누락": sum(r[3] in (None, "") for r in data),
        "일자 형식 오류": sum(not isinstance(r[1], dt.datetime) for r in data),
        "문자형 숫자 금액": sum(isinstance(r[3], str) and to_num(r[3]) is not None for r in data),
        "정상 음수 환입": sum(num(r[3]) and r[3] < 0 for r in data),
        "정상 0원 취소": sum(num(r[3]) and r[3] == 0 for r in data),
        "10억원 초과": sum(num(r[3]) and r[3] > 1e9 for r in data),
        "거래번호 중복 그룹": sum(c > 1 for c in collections.Counter(r[0] for r in data).values()),
        "전체 행 동일 중복 그룹": sum(c > 1 for c in collections.Counter(data).values()),
    }
    bad = [(k, key[k], v) for k, v in got.items() if k in key and round(key[k]) != round(v)]
    calc = collections.defaultdict(lambda: [0, 0, 0])
    for r in data:
        if isinstance(r[1], dt.datetime) and r[5]:
            c = calc[(r[5], r[1].year)]
            c[0] += 1; c[1] += r[3] if num(r[3]) else 0; c[2] += to_num(r[3]) or 0
    bad += [("집계", k, v, tuple(calc[k])) for k, v in agg.items() if tuple(map(round, calc[k])) != v]
    print(f, "확인 항목", len([k for k in got if k in key]), "개 + 회사·연도 집계", len(agg), "개 →", "모두 일치" if not bad else bad)
    return data, not bad

d500, ok1 = check("거래내역_검증용_500행.xlsx")
d10k, ok2 = check("거래내역_대량검증용_10000행.xlsx")
ok3 = d500 == d10k[:500]
print("500행 = 10000행의 앞 500행:", ok3)
sys.exit(0 if ok1 and ok2 and ok3 else 1)
